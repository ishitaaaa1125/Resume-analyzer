from __future__ import annotations

import os
import tempfile
from typing import Final

import pdfplumber
from docx import Document
from fastapi import HTTPException, UploadFile
from pdfminer.pdfparser import PDFSyntaxError

PDF_MIME_TYPES: Final[set[str]] = {"application/pdf"}
DOCX_MIME_TYPES: Final[set[str]] = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class DocumentProcessingError(Exception):
    """Raised when the uploaded resume cannot be processed."""


async def extract_text_from_resume(resume: UploadFile) -> str:
    """Extract text from a PDF or DOCX resume and reject empty extraction."""
    filename = resume.filename or "unknown"
    content_type = resume.content_type or ""
    print(f"Uploaded filename: {filename}")
    print(f"Uploaded content type: {content_type or 'unknown'}")

    is_pdf = content_type in PDF_MIME_TYPES or filename.lower().endswith(".pdf")
    is_docx = content_type in DOCX_MIME_TYPES or filename.lower().endswith(".docx")

    if is_pdf:
        return await extract_text_from_pdf(resume)
    if is_docx:
        return await extract_text_from_docx(resume)

    raise HTTPException(status_code=400, detail="Resume must be a PDF or DOCX file.")


async def extract_text_from_pdf(resume: UploadFile) -> str:
    """Extract text from the uploaded PDF file using pdfplumber."""
    if not resume:
        raise HTTPException(status_code=400, detail="Resume file is required.")

    temp_path: str | None = None
    try:
        content = await resume.read()
        if not content:
            raise DocumentProcessingError("The uploaded PDF is empty.")

        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as temp_file:
            temp_path = temp_file.name
            temp_file.write(content)

        text_parts: list[str] = []
        has_image_content = False
        with pdfplumber.open(temp_path) as pdf:
            page_count = len(pdf.pages)
            print(f"Page count: {page_count}")
            if page_count == 0:
                raise DocumentProcessingError("No pages found in the uploaded PDF.")

            for page_number, page in enumerate(pdf.pages, start=1):
                page_text = (page.extract_text() or "").strip()
                if page_text:
                    text_parts.append(page_text)
                if page.images:
                    has_image_content = True
                    print(f"Page {page_number} contains embedded images; OCR may be required.")

        extracted_text = "\n\n".join(text_parts).strip()
        _log_extraction_preview(extracted_text)

        if not extracted_text:
            raise HTTPException(
                status_code=400,
                detail="Unable to extract text from the uploaded PDF. Please upload a text-based PDF.",
            )
        if has_image_content and len(extracted_text) < 300:
            raise HTTPException(status_code=400, detail="This resume appears to be scanned. OCR is required.")

        return extracted_text
    except (PDFSyntaxError, ValueError) as exc:
        raise HTTPException(status_code=400, detail="The uploaded PDF is corrupted or unreadable.") from exc
    except DocumentProcessingError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=500, detail="Unexpected error while processing the PDF.") from exc
    finally:
        if temp_path and os.path.exists(temp_path):
            os.remove(temp_path)


async def extract_text_from_docx(resume: UploadFile) -> str:
    """Extract text from the uploaded DOCX resume."""
    temp_path: str | None = None
    try:
        content = await resume.read()
        if not content:
            raise DocumentProcessingError("The uploaded DOCX is empty.")

        with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as temp_file:
            temp_path = temp_file.name
            temp_file.write(content)

        document = Document(temp_path)
        paragraph_text = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        table_text: list[str] = []
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    table_text.append(" | ".join(cells))

        extracted_text = "\n".join(paragraph_text + table_text).strip()
        print("Page count: DOCX")
        _log_extraction_preview(extracted_text)

        if not extracted_text:
            raise HTTPException(
                status_code=400,
                detail="Unable to extract text from the uploaded DOCX. Please upload a text-based resume.",
            )

        return extracted_text
    except DocumentProcessingError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail="The uploaded DOCX is corrupted or unreadable.") from exc
    finally:
        if temp_path and os.path.exists(temp_path):
            os.remove(temp_path)


def _log_extraction_preview(extracted_text: str) -> None:
    print(f"Extracted character count: {len(extracted_text)}")
    print("First 1500 characters:")
    print(extracted_text[:1500])
